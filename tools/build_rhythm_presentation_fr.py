# -*- coding: utf-8 -*-
"""
Génère une présentation FR sur le mini-jeu de rythme (CookingGame.cs),
style proche des exemples (fond clair, carte arrondie, blocs code sombres avec numéros de ligne).
Réutilise des captures extraites de « presentation groupe sans video.pptx ».
"""
from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE_PPTX = ROOT / "presentation groupe sans video.pptx"
SOURCE_CODE_PATH = ROOT / "Assets" / "Scenes" / "script" / "maison" / "CookingGame.cs"
OUTPUT_PPTX = ROOT / "Presentation_MiniJeu_Rythme_FR_v3.pptx"

# Couleurs (fond « normal » + thème sombre code)
BG_SLIDE = RGBColor(255, 255, 255)
CARD_LINE = RGBColor(0xD8, 0xD8, 0xE0)
CODE_BG = (0x28, 0x2C, 0x34)
CODE_GUTTER = (0x31, 0x36, 0x3F)
CODE_LINENO = (0x7C, 0x81, 0x8E)
CODE_TEXT = (0xDC, 0xDC, 0xE0)


def _find_mono_font(size: int) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\consola.ttf"),
        Path(r"C:\Windows\Fonts\cour.ttf"),
        Path(r"C:\Windows\Fonts\lucon.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def render_code_image(
    lines: list[tuple[int, str]],
    filename_caption: str,
    out_path: Path,
    font_size: int = 19,
    line_h: int = 26,
    pad: int = 18,
) -> None:
    font = _find_mono_font(font_size)
    max_w = 0
    for _, ln in lines:
        bbox = font.getbbox(ln.replace("\t", "    "))
        max_w = max(max_w, bbox[2] - bbox[0])
    gutter_w = 56
    img_w = pad * 2 + gutter_w + max_w + 40
    img_h = pad * 2 + len(lines) * line_h + 50
    im = Image.new("RGB", (img_w, img_h), CODE_BG)
    draw = ImageDraw.Draw(im)
    y = pad
    lineno_font = _find_mono_font(font_size - 1)
    for num, text in lines:
        draw.text((pad, y), f"{num:4d}", fill=CODE_LINENO, font=lineno_font)
        draw.rectangle([pad + gutter_w - 8, y, pad + gutter_w - 2, y + line_h - 4], fill=CODE_GUTTER)
        draw.text((pad + gutter_w + 6, y), text.replace("\t", "    "), fill=CODE_TEXT, font=font)
        y += line_h
    cap = f"Capture : {filename_caption}"
    cap_font = _find_mono_font(15)
    draw.text((pad, img_h - 36), cap, fill=CODE_LINENO, font=cap_font)
    im.save(out_path, "PNG")


def extract_slide_pictures(prs: Presentation, slide_idx: int, dest_dir: Path) -> list[Path]:
    """slide_idx: 0-based. Retourne les fichiers PNG/JPEG écrits."""
    dest_dir.mkdir(parents=True, exist_ok=True)
    out: list[Path] = []
    slide = prs.slides[slide_idx]
    for i, sh in enumerate(slide.shapes):
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        ext = sh.image.ext
        path = dest_dir / f"src_slide{slide_idx + 1}_img{len(out) + 1}.{ext}"
        path.write_bytes(sh.image.blob)
        out.append(path)
    return out


def set_slide_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rounded_card(slide, prs: Presentation):
    m = Inches(0.42)
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        m,
        m,
        prs.slide_width - 2 * m,
        prs.slide_height - 2 * m,
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = CARD_LINE
    card.line.width = Pt(1)
    try:
        card.adjustments[0] = 0.06
    except Exception:
        pass
    # envoyer la carte au fond
    sp = slide.shapes._spTree
    sp.remove(card._element)
    sp.insert(2, card._element)
    return card


def add_title_text(slide, left, top, width, height, text: str, size_pt: int = 30, bold: bool = True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.name = "Calibri Light"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x2E)
    return box


def add_body_bullets(slide, left, top, width, height, lines: list[str], size_pt: int = 15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(size_pt)
        p.font.name = "Calibri"
        p.font.color.rgb = RGBColor(0x33, 0x33, 0x40)
        p.space_after = Pt(6)
    return box


def add_picture_fit(slide, path: Path, left, top, max_w, max_h):
    """max_w, max_h : objets Length (ex. Inches(6))."""
    im = Image.open(path)
    w_px, h_px = im.size
    max_w_in = max_w.inches
    max_h_in = max_h.inches
    # ~96 px par pouce côté PowerPoint
    ratio = min(max_w_in / (w_px / 96.0), max_h_in / (h_px / 96.0))
    slide.shapes.add_picture(
        str(path),
        left,
        top,
        width=Inches((w_px / 96.0) * ratio),
        height=Inches((h_px / 96.0) * ratio),
    )


def main():
    if not SOURCE_PPTX.exists():
        raise SystemExit(f"Fichier source introuvable : {SOURCE_PPTX}")

    tmp = Path(tempfile.mkdtemp(prefix="ppt_rhythm_"))
    code_dir = tmp / "code_png"
    code_dir.mkdir()

    src_prs = Presentation(str(SOURCE_PPTX))
    # Slides source : 61 = mécanique + 2 captures
    imgs_61 = extract_slide_pictures(src_prs, 60, tmp / "s61")

    prs = Presentation()
    prs.slide_width = int(13.333333 * 914400)  # 16:9
    prs.slide_height = int(7.5 * 914400)
    blank = prs.slide_layouts[6]

    # --- Slide titre ---
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide, BG_SLIDE)
    add_title_text(
        slide,
        Inches(0.85),
        Inches(2.2),
        Inches(11.5),
        Inches(1.2),
        "DUO SENSE",
        40,
        True,
    )
    add_title_text(
        slide,
        Inches(0.85),
        Inches(3.1),
        Inches(11.5),
        Inches(1),
        "Mini-jeu de rythme (cuisine)",
        28,
        False,
    )
    add_body_bullets(
        slide,
        Inches(0.85),
        Inches(4.2),
        Inches(11.5),
        Inches(2),
        [
            "Script : CookingGame.cs — écran partagé coopératif",
            "Iris : rythme audio (fenêtre de frappe) · Achille : notes tombantes",
        ],
        18,
    )

    # --- Contexte & but ---
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide, BG_SLIDE)
    add_title_text(slide, Inches(0.75), Inches(0.55), Inches(12), Inches(0.7), "Contexte & but", 32)
    add_body_bullets(
        slide,
        Inches(0.75),
        Inches(1.25),
        Inches(12),
        Inches(1.1),
        [
            "Rôle : mini-jeu d’interaction (IInteractable) intégré à la quête cuisine.",
            "Responsabilité : boucle de jeu, entrées clavier, feedback UI et audio fiable.",
        ],
        15,
    )
    gap = Inches(0.25)
    usable_w = Inches(12.35)
    pic_w = Inches(usable_w.inches / 2 - gap.inches / 2)
    top_img = Inches(2.35)
    max_h = Inches(4.35)
    left0 = Inches(0.75)
    if len(imgs_61) >= 1:
        add_picture_fit(slide, imgs_61[0], left0, top_img, pic_w, max_h)
    if len(imgs_61) >= 2:
        add_picture_fit(slide, imgs_61[1], left0 + pic_w + gap, top_img, pic_w, max_h)

    # --- Qualité / diffusion (Inspector) ---
    slide = prs.slides.add_slide(blank)
    set_slide_background(slide, BG_SLIDE)
    add_title_text(slide, Inches(0.75), Inches(0.55), Inches(12), Inches(0.7), "Paramètres (Inspector)", 32)
    add_body_bullets(
        slide,
        Inches(0.75),
        Inches(1.2),
        Inches(5.6),
        Inches(1.6),
        [
            "Paramètres exposés : durée, intervalle de beat, vitesse des notes, fenêtre de hit.",
            "Réglage sans recompiler — itération design rapide.",
            "Résultat : le mini-jeu reste ajustable pour la quête.",
        ],
        14,
    )

    # --- Code : extraits fidèles de CookingGame.cs (sans lignes en cyrillique) ---
    if not SOURCE_CODE_PATH.exists():
        raise SystemExit(f"Fichier source introuvable : {SOURCE_CODE_PATH}")

    code_text = SOURCE_CODE_PATH.read_text(encoding="utf-8-sig", errors="ignore")
    code_lines = code_text.splitlines()

    def has_cyrillic(s: str) -> bool:
        # On retire toute ligne contenant des caractères cyrilliques (russe/ukrainien/...)
        return re.search(r"[\u0400-\u04FF]", s or "") is not None

    def extract_numbered_lines(start: int, end: int) -> list[tuple[int, str]]:
        out: list[tuple[int, str]] = []
        for ln in range(start, end + 1):
            if ln < 1 or ln > len(code_lines):
                continue
            raw = code_lines[ln - 1]
            if has_cyrillic(raw):
                continue
            out.append((ln, raw.rstrip()))
        # Si tout a été filtré, on évite les slides vides
        return out

    # Plages de lignes (1-indexées) dans CookingGame.cs
    snippets = [
        {
            "title": "Données de note (FallingNote)",
            "caption": "CookingGame.cs · FallingNote",
            "start": 828,
            "end": 836,
            "bullets": [
                "Chaque note transporte seulement l’info UI + la colonne (gauche/droit) + l’état de frappe.",
                "Résultat : la mise à jour et la validation restent simples.",
            ],
        },
        {
            "title": "Boucle Update()",
            "caption": "CookingGame.cs · Update()",
            "start": 126,
            "end": 144,
            "bullets": [
                "Fin de partie centralisée (temps ou vies), puis mise à jour : rythme (Iris) et notes (Achille).",
                "Résultat : même cœur de jeu, pas de duplication.",
            ],
        },
        {
            "title": "Iris : beat + fenêtre de frappe",
            "caption": "CookingGame.cs · StartRhythmBeat()",
            "start": 194,
            "end": 221,
            "bullets": [
                "Active la fenêtre de réponse (hitWindow) et place le son via listener / caméra.",
                "Résultat : le beat reste audible et visuel synchronisé.",
            ],
        },
        {
            "title": "Achille : spawn + validation",
            "caption": "CookingGame.cs · SpawnNote() / TryHitNote()",
            "start": 322,
            "end": 384,
            "bullets": [
                "Spawn aléatoire sur 2 colonnes (gauche/droit) et validation de la note « hittable ».",
                "Résultat : hit/miss immédiat + destruction de la note validée.",
            ],
        },
        {
            "title": "Lien avec la quête cuisine",
            "caption": "CookingGame.cs · CompleteCookingQuestStep()",
            "start": 780,
            "end": 823,
            "bullets": [
                "Le mini-jeu trouve la quête active et marque l’étape 2 comme complétée.",
                "Si la quête est terminée, elle est retirée du système.",
            ],
        },
    ]

    for si, sn in enumerate(snippets):
        numbered_lines = extract_numbered_lines(sn["start"], sn["end"])
        if not numbered_lines:
            continue
        png = code_dir / f"code_{sn['start']}_{si}.png"
        render_code_image(numbered_lines, sn["caption"], png)
        slide = prs.slides.add_slide(blank)
        set_slide_background(slide, BG_SLIDE)
        add_title_text(
            slide,
            Inches(0.75),
            Inches(0.5),
            Inches(12),
            Inches(0.65),
            sn["title"],
            26,
        )
        pic_left = Inches(0.65)
        pic_top = Inches(1.15)
        pic_max_w = Inches(12.1)
        pic_max_h = Inches(5.35)
        add_picture_fit(slide, png, pic_left, pic_top, pic_max_w, pic_max_h)
        add_body_bullets(
            slide,
            Inches(0.75),
            Inches(6.55),
            Inches(11.8),
            Inches(0.85),
            sn["bullets"],
            13,
        )

    prs.save(str(OUTPUT_PPTX))
    print(f"OK -> {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()
