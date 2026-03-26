# -*- coding: utf-8 -*-
"""
Présentation FR 'jury-ready' (8 slides) pour le mini-jeu de rythme.

Design goals:
- pas de grosses cartes / pas de "половина пустая"
- code lisible: grand bloc, inséré pour remplir la zone
- pas de russe dans les captures de code (sanitise les lignes en cyrillique)
- structure: Problème -> Solution -> Iris -> Achille -> Fin -> Quête -> Conclusion
"""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


ROOT = Path(__file__).resolve().parent.parent
CODE_PATH = ROOT / "Assets" / "Scenes" / "script" / "maison" / "CookingGame.cs"
OUTPUT_PPTX = ROOT / "Presentation_MiniJeu_Rythme_FR_jury2.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN_X = 0.75

CODE_LEFT = Inches(MARGIN_X)
CODE_TOP = Inches(1.55)
CODE_W = Inches(SLIDE_W - 2 * MARGIN_X)
CODE_H = Inches(4.55)


SLIDE_BG = RGBColor(0xF6, 0xF7, 0xFB)
TITLE_COLOR = RGBColor(0x1E, 0x1E, 0x2A)
SUBTITLE_COLOR = RGBColor(0x4B, 0x4F, 0x5F)
BODY_COLOR = RGBColor(0x2C, 0x2C, 0x38)

ACCENT = RGBColor(0x49, 0x7B, 0xFF)
CARD_LINE = RGBColor(0xE3, 0xE6, 0xEE)

CODE_BG = (0x1E, 0x24, 0x2F)
CODE_GUTTER = (0x2A, 0x31, 0x44)
CODE_LINENO = (0x9CA3AF)
CODE_TEXT = (0xE5, 0xE7, 0xEB)


def _find_mono_font(size: int) -> ImageFont.FreeTypeFont:
    candidates = [
        Path(r"C:\Windows\Fonts\consola.ttf"),
        Path(r"C:\Windows\Fonts\cour.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size=size)
    return ImageFont.load_default()


def _sanitize_cyrillic(line: str) -> str:
    # Remove cyrillic chars => replace with spaces, so no Russian letters appear.
    if not line:
        return line
    line = re.sub(r"[\u0400-\u04FF]", " ", line)
    line = re.sub(r"[ \t]+", " ", line)
    return line.rstrip("\n")


def _wrap_code(s: str, max_chars: int = 88) -> list[str]:
    s = s.rstrip()
    if len(s) <= max_chars:
        return [s]
    out = []
    cur = s
    while len(cur) > max_chars:
        out.append(cur[: max_chars - 3] + "...")
        cur = cur[max_chars - 3 :]
    if cur:
        out.append(cur)
    return out


def render_code_png(lines: list[tuple[int, str]], out_png: Path) -> None:
    """
    lines: [(line_no, text)]
    Generate a PNG sized for good readability (not too tall).
    """
    font_size = 16
    line_h = 21
    pad_x = 14
    pad_y = 10
    gutter_w = 62

    font = _find_mono_font(font_size)
    lineno_font = _find_mono_font(max(12, font_size - 2))

    # Prepare wrapped content
    wrapped: list[tuple[int, str]] = []
    for ln, raw in lines:
        raw = _sanitize_cyrillic(raw)
        parts = _wrap_code(raw, max_chars=92)
        for i, p in enumerate(parts):
            wrapped.append((ln if i == 0 else -1, p))

    # Estimate width by longest line
    max_text_w = 0
    for _, t in wrapped:
        bbox = font.getbbox(t if t else " ")
        max_text_w = max(max_text_w, bbox[2] - bbox[0])

    img_w = pad_x * 2 + gutter_w + max_text_w + 10
    img_h = pad_y * 2 + len(wrapped) * line_h + 24

    im = Image.new("RGB", (img_w, img_h), CODE_BG)
    draw = ImageDraw.Draw(im)

    y = pad_y
    for ln, text in wrapped:
        # gutter tag
        draw.rectangle(
            [pad_x + gutter_w - 10, y + 3, pad_x + gutter_w - 2, y + line_h - 6],
            fill=CODE_GUTTER,
        )
        if ln != -1:
            draw.text((pad_x + 6, y), f"{ln:4d}", fill=CODE_LINENO, font=lineno_font)
        else:
            draw.text((pad_x + 6, y), "    ", fill=CODE_LINENO, font=lineno_font)

        draw.text((pad_x + gutter_w + 6, y), text, fill=CODE_TEXT, font=font)
        y += line_h

    out_png.parent.mkdir(parents=True, exist_ok=True)
    im.save(out_png, "PNG")


def add_slide_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = SLIDE_BG


def add_title(slide, text: str, top_in: float = 0.55, font_size: int = 34) -> None:
    box = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(top_in), Inches(SLIDE_W - 2 * MARGIN_X), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.name = "Calibri"
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR


def add_subtitle(slide, text: str, top_in: float = 1.15) -> None:
    box = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(top_in), Inches(SLIDE_W - 2 * MARGIN_X), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.name = "Calibri"
    p.font.color.rgb = SUBTITLE_COLOR


def add_bullets(slide, lines: list[str], left_in: float, top_in: float, width_in: float, height_in: float, font_size: int = 16) -> None:
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    tf = box.text_frame
    tf.clear()
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.name = "Calibri"
        p.font.color.rgb = BODY_COLOR
        p.space_after = Pt(4)


def extract_code_lines(code_text: list[str], line_nums: list[int]) -> list[tuple[int, str]]:
    out = []
    for ln in line_nums:
        if 1 <= ln <= len(code_text):
            out.append((ln, code_text[ln - 1]))
    return out


def main() -> None:
    if not CODE_PATH.exists():
        raise SystemExit(f"Fichier introuvable: {CODE_PATH}")

    code_lines = CODE_PATH.read_text(encoding="utf-8-sig", errors="ignore").splitlines()

    tmp_dir = ROOT / "_tmp_jury_code_imgs"
    tmp_dir.mkdir(exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # Shared style: thin underline accent under title (rectangle)
    def add_underline(slide, y_in: float) -> None:
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN_X), Inches(y_in), Inches(4.2), Inches(0.12))
        r.fill.solid()
        r.fill.fore_color.rgb = ACCENT
        r.line.width = Pt(0)

    # 1) Intro
    s = prs.slides.add_slide(blank)
    add_slide_bg(s)
    add_title(s, "Mini-jeu de rythme (cuisine)", 0.45, 36)
    add_underline(s, 1.28)
    add_bullets(
        s,
        [
            "Iris: fenêtre temporelle + feedback visuel/son.",
            "Achille: spawn de notes + validation par timing.",
            "Résultat: hit/miss immédiat + fin claire.",
        ],
        MARGIN_X,
        1.65,
        SLIDE_W - 2 * MARGIN_X,
        5.5,
        18,
    )

    # 2) Problème initial
    s2 = prs.slides.add_slide(blank)
    add_slide_bg(s2)
    add_title(s2, "Problème initial", 0.45, 36)
    add_underline(s2, 1.28)
    add_bullets(
        s2,
        [
            "Le rythme doit être lisible: QUAND appuyer (hitWindow).",
            "En parallèle, les notes doivent tomber et être validées seulement au bon moment.",
            "Attendu: feedback immédiat et boucle de fin déterministe.",
        ],
        MARGIN_X,
        1.65,
        SLIDE_W - 2 * MARGIN_X,
        5.5,
        18,
    )

    # 3) Solution
    s3 = prs.slides.add_slide(blank)
    add_slide_bg(s3)
    add_title(s3, "Solution", 0.45, 36)
    add_underline(s3, 1.28)
    # 2x2 grid
    grid_left = MARGIN_X
    grid_top = 1.65
    box_w = (SLIDE_W - 2 * MARGIN_X - 0.35) / 2
    box_h = 1.4
    gap_x = 0.35
    gap_y = 0.35
    titles = [
        "Iris (rythme)\n+ fenêtre",
        "Achille (notes)\n+ validation",
        "Boucle unique\nUpdate()",
        "Fin unique\nEndGame()",
    ]
    positions = [
        (grid_left, grid_top),
        (grid_left + box_w + gap_x, grid_top),
        (grid_left, grid_top + box_h + gap_y),
        (grid_left + box_w + gap_x, grid_top + box_h + gap_y),
    ]
    for (x_in, y_in), t in zip(positions, titles):
        r = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_in), Inches(y_in), Inches(box_w), Inches(box_h))
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor(255, 255, 255)
        r.line.color.rgb = CARD_LINE
        r.line.width = Pt(1)
        r.text_frame.text = t
        # Force a readable size
        tf = r.text_frame
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.name = "Calibri"
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = 1  # center

    add_bullets(
        s3,
        ["Résultat: logique séparée mais synchronisée sur le même temps de jeu."],
        MARGIN_X + 0.1,
        grid_top + (box_h + gap_y) * 2 + 0.25,
        SLIDE_W - 2 * MARGIN_X - 0.2,
        1.2,
        16,
    )

    # Code slides helper
    def add_code_slide(
        title_text: str,
        code_title: str,
        bullet_lines: list[str],
        line_nums: list[int],
    ) -> None:
        nonlocal prs
        s = prs.slides.add_slide(blank)
        add_slide_bg(s)
        add_title(s, title_text, 0.45, 32)
        add_underline(s, 1.23)
        add_subtitle(s, code_title, 1.25)

        extracted = extract_code_lines(code_lines, line_nums)
        png = tmp_dir / f"code_{title_text.replace(' ', '_')}.png"
        render_code_png(extracted, png)
        # Insert so it fills code area (scale/crop handled by PPT scaling)
        s.shapes.add_picture(str(png), CODE_LEFT, CODE_TOP, width=CODE_W, height=CODE_H)

        add_bullets(
            s,
            bullet_lines,
            MARGIN_X,
            6.15,
            SLIDE_W - 2 * MARGIN_X,
            1.2,
            16,
        )

    # 4) Iris
    iris_nums = [
        146, 151, 152, 153, 154, 157, 158, 160, 162, 166, 170, 171, 175,
        179, 180, 182, 183, 185, 189, 190,
        194, 196, 197, 199, 202, 204, 206, 207, 208, 209, 210, 212, 213, 214, 217, 219,
    ]
    add_code_slide(
        "Iris: fenêtre de frappe (son → hit)",
        "CookingGame.cs · StartRhythmBeat + fenêtre",
        [
            "beatTimer compare beatInterval → StartRhythmBeat().",
            "canWomanPress + womanPressWindow décident du hit/miss.",
            "Feedback: rythme UI via rhythmIndicator.",
        ],
        iris_nums,
    )

    # 5) Achille
    ach_nums = [
        322, 323, 324, 327, 328, 332,
        349, 350, 353, 354,
        356, 357, 360, 361, 363, 364, 365, 366,
        359, 380, 383, 384,
    ]
    add_code_slide(
        "Achille: spawn + validation des notes",
        "CookingGame.cs · SpawnNote + TryHitNote",
        [
            "Spawn aléatoire: 2 colonnes → activeNotes.",
            "TryHitNote ne valide que note.canBeHit.",
            "Résultat: ManHit()/ManMiss() + destruction de la note validée.",
        ],
        ach_nums,
    )

    # 6) Fin
    end_nums = [
        126, 130, 132, 133, 135, 136, 141, 142, 143, 144,
        656, 658, 661,
        663, 665, 667, 669, 670,
        679, 684,
        686, 688, 689, 691, 692,
    ]
    add_code_slide(
        "Fin de partie: conditions + feedback",
        "CookingGame.cs · Update() + EndGame()",
        [
            "Stop unique: temps écoulé ou vies à 0 → EndGame().",
            "Succès: panneau success + Bravo! + CompleteCookingQuestStep().",
            "Échec: panneau GameOver + Perdu!",
        ],
        end_nums,
    )

    # 7) Lien quête
    quest_nums = [
        782, 784, 785, 789,
        791, 793, 795, 797, 800,
        802, 806,
        809, 810, 813,
        815, 816,
        820, 822, 823,
    ]
    add_code_slide(
        "Lien avec la quête cuisine",
        "CookingGame.cs · CompleteCookingQuestStep()",
        [
            "Cherche la quête active par questName.",
            "Complète l’étape 2: foundQuest.CompleteStep(1).",
            "Si la quête est finie: CompleteQuest(foundQuest).",
        ],
        quest_nums,
    )

    # 8) Conclusion
    s8 = prs.slides.add_slide(blank)
    add_slide_bg(s8)
    add_title(s8, "Conclusion", 0.45, 36)
    add_underline(s8, 1.28)
    add_bullets(
        s8,
        [
            "Découpage clair: Iris gère le timing, Achille gère les notes.",
            "Boucle stable: un seul Update() et une fin unique EndGame().",
            "Intégration: succès du mini-jeu → progression de quête.",
        ],
        MARGIN_X,
        1.65,
        SLIDE_W - 2 * MARGIN_X,
        5.5,
        18,
    )

    prs.save(str(OUTPUT_PPTX))
    print(f"OK -> {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()

