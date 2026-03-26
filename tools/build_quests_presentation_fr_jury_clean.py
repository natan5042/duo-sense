# -*- coding: utf-8 -*-
"""
Présentation FR (8 slides) pour la partie "Quêtes" (QuestSystem).

Design:
- fond clair simple
- pas de grosses cartes pleine page
- code lisible: gros bloc fixe + bullets courts
- pas de russe dans les captures de code (sanitization des lignes en cyrillique)
"""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT_PPTX = ROOT / "Presentation_Quests_FR_jury_clean_v2.pptx"

QUEST_SYSTEM_CS = ROOT / "Assets" / "Scripts" / "QuestSystem.cs"
QUEST_UI_CS = ROOT / "Assets" / "Scripts" / "QuestUI.cs"
DOOR_UNLOCK_CS = ROOT / "Assets" / "Scripts" / "DoorUnlockOnQuestsComplete.cs"
SUPERMARKET_MANAGER_CS = ROOT / "Assets" / "Scenes" / "script" / "niveau_supermarché" / "SupermarketQuestManager.cs"

for p in [QUEST_SYSTEM_CS, QUEST_UI_CS, DOOR_UNLOCK_CS, SUPERMARKET_MANAGER_CS]:
    if not p.exists():
        raise SystemExit(f"Fichier introuvable: {p}")


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_X = 0.75

# Colors
SLIDE_BG = RGBColor(0xF6, 0xF7, 0xFB)
TITLE_COLOR = RGBColor(0x1E, 0x1E, 0x2A)
BODY_COLOR = RGBColor(0x2C, 0x2C, 0x38)
SUBTITLE_COLOR = RGBColor(0x4B, 0x4F, 0x5F)
ACCENT = RGBColor(0x49, 0x7B, 0xFF)
CARD_LINE = RGBColor(0xE3, 0xE6, 0xEE)

# Code theme
CODE_BG = (0x1E, 0x24, 0x2F)
CODE_GUTTER = (0x2A, 0x31, 0x44)
CODE_LINENO = (0x9C, 0xA3, 0xAF)
CODE_TEXT = (0xE5, 0xE7, 0xEB)


def _find_mono_font(size: int) -> ImageFont.FreeTypeFont:
    # Consolas tends to look clean.
    candidates = [
        Path(r"C:\Windows\Fonts\consola.ttf"),
        Path(r"C:\Windows\Fonts\cour.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size=size)
    return ImageFont.load_default()


def _sanitize_cyrillic(line: str) -> str:
    # Replace any cyrillic char with whitespace (keeps braces/structure visible).
    line = re.sub(r"[\u0400-\u04FF]", " ", line or "")
    line = re.sub(r"[ \t]+", " ", line)
    return line.rstrip("\n")


def _wrap_code_line(s: str, max_chars: int = 92) -> list[str]:
    s = s.rstrip()
    if len(s) <= max_chars:
        return [s]
    out: list[str] = []
    cur = s
    while len(cur) > max_chars:
        out.append(cur[: max_chars - 3] + "...")
        cur = cur[max_chars - 3 :]
    if cur:
        out.append(cur)
    return out


def render_code_png(
    lines: list[tuple[int, str]],
    out_png: Path,
    *,
    font_size: int = 16,
    line_h: int = 21,
) -> None:
    out_png.parent.mkdir(parents=True, exist_ok=True)
    font = _find_mono_font(font_size)
    lineno_font = _find_mono_font(max(12, font_size - 2))

    # Prepare wrapped lines with line numbers only on the first wrapped segment.
    wrapped: list[tuple[int, str]] = []
    for ln, raw in lines:
        text = _sanitize_cyrillic(raw)
        # Large max_chars => fewer wraps => shorter PNG => less scaling (no empty zones).
        parts = _wrap_code_line(text, max_chars=130)
        for i, p in enumerate(parts):
            wrapped.append((ln if i == 0 else -1, p))

    # Compute width from the longest part.
    max_w = 0
    for _, t in wrapped:
        bbox = font.getbbox(t if t else " ")
        max_w = max(max_w, bbox[2] - bbox[0])

    pad_x = 14
    pad_y = 10
    gutter_w = 62

    img_w = pad_x * 2 + gutter_w + max_w + 10
    img_h = pad_y * 2 + len(wrapped) * line_h + 18

    im = Image.new("RGB", (img_w, img_h), CODE_BG)
    draw = ImageDraw.Draw(im)

    y = pad_y
    for ln, text in wrapped:
        # gutter tag
        draw.rectangle(
            [pad_x + gutter_w - 10, y + 3, pad_x + gutter_w - 2, y + line_h - 6],
            fill=CODE_GUTTER,
        )
        if ln != -1:
            draw.text((pad_x + 6, y), f"{ln:4d}", fill=CODE_LINENO, font=lineno_font)
        else:
            draw.text((pad_x + 6, y), "    ", fill=CODE_LINENO, font=lineno_font)
        draw.text((pad_x + gutter_w + 6, y), text, fill=CODE_TEXT, font=font)
        y += line_h

    im.save(out_png, "PNG")


def add_slide_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = SLIDE_BG


def add_title(slide, text: str, y_in: float, size: int = 34) -> None:
    box = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(y_in), SLIDE_W - Inches(2 * MARGIN_X), Inches(0.6))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.name = "Calibri"
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR


def add_underline(slide, y_in: float) -> None:
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN_X), Inches(y_in), Inches(4.2), Inches(0.12))
    r.fill.solid()
    r.fill.fore_color.rgb = ACCENT
    r.line.width = Pt(0)


def add_bullets(slide, lines: list[str], left_in: float, top_in: float, width_in: float, height_in: float, font_size: int = 16) -> None:
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.name = "Calibri"
        p.font.color.rgb = BODY_COLOR
        p.space_after = Pt(4)


def add_picture_fit(slide, img_path: Path, left_in: float, top_in: float, max_w_in: float, max_h_in: float) -> None:
    im = Image.open(img_path)
    w_px, h_px = im.size
    # 96 px ~ 1 inch
    max_w_px = max_w_in * 96.0
    max_h_px = max_h_in * 96.0
    scale = min(max_w_px / max(w_px, 1), max_h_px / max(h_px, 1))
    out_w_in = (w_px * scale) / 96.0
    out_h_in = (h_px * scale) / 96.0
    slide.shapes.add_picture(str(img_path), Inches(left_in), Inches(top_in), width=Inches(out_w_in), height=Inches(out_h_in))


def read_lines(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8-sig", errors="ignore").splitlines()


def extract_ranges(code_lines: list[str], ranges: list[tuple[int, int]]) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for start, end in ranges:
        for ln in range(start, end + 1):
            if 1 <= ln <= len(code_lines):
                out.append((ln, code_lines[ln - 1]))
    return out


def slug(s: str) -> str:
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s or "code"


def main() -> None:
    qsys_lines = read_lines(QUEST_SYSTEM_CS)
    qui_lines = read_lines(QUEST_UI_CS)
    door_lines = read_lines(DOOR_UNLOCK_CS)
    super_lines = read_lines(SUPERMARKET_MANAGER_CS)

    tmp_dir = ROOT / "_tmp_quests_code_imgs"
    tmp_dir.mkdir(exist_ok=True)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # Common code area (fill most of the slide)
    code_left = 0.75
    code_top = 1.5
    code_w = 11.83
    code_h = 4.75

    # --- Slide 1: Title/Context ---
    s = prs.slides.add_slide(blank)
    add_slide_bg(s)
    add_title(s, "Système de quêtes (Unity)", 0.55, 36)
    add_underline(s, 1.22)
    add_bullets(
        s,
        [
            "Objectif: une progression de quêtes stable entre scènes.",
            "Approche: un état central (QuestSystem) + UI réactive via événement.",
            "Résultat attendu: pas de polling lourd, logique claire et réutilisable.",
        ],
        left_in=MARGIN_X,
        top_in=1.6,
        width_in=float(13.333 - 2 * MARGIN_X),
        height_in=5.0,
        font_size=18,
    )

    # --- Slide 2: Problème initial ---
    s2 = prs.slides.add_slide(blank)
    add_slide_bg(s2)
    add_title(s2, "Problème initial", 0.55, 36)
    add_underline(s2, 1.22)
    add_bullets(
        s2,
        [
            "Les quêtes doivent persister et rester cohérentes quand on change de scène.",
            "L’UI doit refléter immédiatement l’avancement (sinon: confusion pour le joueur).",
            "Le gameplay doit savoir quand une quête est finie (ex: porte à cacher).",
        ],
        left_in=MARGIN_X,
        top_in=1.6,
        width_in=13.333 - 2 * MARGIN_X,
        height_in=5.0,
        font_size=18,
    )

    # --- Slide 3: Solution (architecture) ---
    s3 = prs.slides.add_slide(blank)
    add_slide_bg(s3)
    add_title(s3, "Solution", 0.55, 36)
    add_underline(s3, 1.22)

    # 2x2 outlined boxes
    def box(x, y, w, h, text):
        r = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid()
        r.fill.fore_color.rgb = SLIDE_BG
        r.line.color.rgb = CARD_LINE
        r.line.width = Pt(1)
        tf = r.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(16)
        p.font.name = "Calibri"
        p.font.bold = True
        tf.word_wrap = True
        return r

    bw = 5.6
    bh = 1.35
    gapx = 0.65
    gapy = 0.35
    x0 = MARGIN_X
    y0 = 1.65
    box(x0, y0, bw, bh, "Quest\n(texte + étapes)")
    box(x0 + bw + gapx, y0, bw, bh, "QuestSystem\nétat global + événement")
    box(x0, y0 + bh + gapy, bw, bh, "QuestUI\naffiche les étapes")
    box(x0 + bw + gapx, y0 + bh + gapy, bw, bh, "Consommateurs\nporte / supermarché")

    add_bullets(
        s3,
        [
            "Un seul point de vérité: `QuestSystem.activeQuests`.",
            "Quand une étape change: `NotifyQuestsChanged()` → UI + logique monde.",
        ],
        left_in=MARGIN_X,
        top_in=4.3,
        width_in=13.333 - 2 * MARGIN_X,
        height_in=2.4,
        font_size=18,
    )

    # --- Helper to create a code slide ---
    def add_code_slide(title, subtitle, code_img_lines, bullets, filename_slug):
        s = prs.slides.add_slide(blank)
        add_slide_bg(s)
        add_title(s, title, 0.55, 30)
        add_underline(s, 1.22)
        # subtitle
        sub_box = s.shapes.add_textbox(Inches(MARGIN_X), Inches(1.35), Inches(11.9), Inches(0.5))
        tf = sub_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(16)
        p.font.name = "Calibri"
        p.font.color.rgb = SUBTITLE_COLOR
        p.font.bold = False

        img = tmp_dir / f"code_{filename_slug}.png"
        render_code_png(code_img_lines, img)
        add_picture_fit(s, img, code_left, code_top, code_w, code_h)

        add_bullets(
            s,
            bullets,
            left_in=MARGIN_X,
            top_in=6.3,
            width_in=13.333 - 2 * MARGIN_X,
            height_in=1.1,
            font_size=16,
        )

    # Slide 4: Quest model (Quest)
    quest_model_lines = []
    # Lines chosen to keep PNG height stable (avoid "empty width" scaling).
    for ln in [17, 18, 21, 22, 39, 40, 41, 42, 43, 44, 45, 46, 73, 75, 76, 80, 82, 84]:
        if 1 <= ln <= len(qsys_lines):
            quest_model_lines.append((ln, qsys_lines[ln - 1]))
    add_code_slide(
        "Modèle de quête",
        "QuestSystem.cs · class Quest",
        quest_model_lines,
        [
            "Les étapes = `steps[]`, l’état = `completedSteps[]`.",
            "Support quantité: `stepRequiredCounts[]` + `stepCurrentCounts[]`.",
            "Quand une étape est complétée: notification UI.",
        ],
        "quest_model",
    )

    # Slide 5: QuestSystem (singleton + événement)
    quest_system_lines = []
    for ln in [96, 98, 106, 108, 110, 111, 130, 132, 133, 139, 140, 145, 159, 163, 198, 200, 202, 204]:
        if 1 <= ln <= len(qsys_lines):
            quest_system_lines.append((ln, qsys_lines[ln - 1]))
    add_code_slide(
        "Gestionnaire global (QuestSystem)",
        "QuestSystem.cs · état + événement",
        quest_system_lines,
        [
            "Singleton `DontDestroyOnLoad` pour persister entre scènes.",
            "Reset supermarché: `RemoveQuestsByTitles()` lors du changement de scène.",
            "Point clé: `OnQuestsChanged` (pas de polling).",
        ],
        "quest_system",
    )

    # Slide 6: UI consumer (QuestUI)
    quest_ui_lines = []
    for ln in [239, 241, 250, 251, 252, 253, 254, 262, 266, 268, 270, 271, 272, 273, 277, 284, 286, 290, 297]:
        if 1 <= ln <= len(qui_lines):
            quest_ui_lines.append((ln, qui_lines[ln - 1]))
    add_code_slide(
        "UI réactive (QuestUI)",
        "QuestUI.cs · UpdateQuestDisplay()",
        quest_ui_lines,
        [
            "Lit `QuestSystem.activeQuests` et reconstruit le texte.",
            "Affiche ✓/○ selon `completedSteps`.",
            "Affiche aussi les quantités (cur/req) quand nécessaire.",
        ],
        "quest_ui",
    )

    # Slide 7: Consommateur monde (porte)
    # Choix: Start subscription + CheckQuestsCompletion + HideDoor
    door_lines_sel = []
    for ln in [54, 56, 59, 61, 86, 88, 90, 101, 103, 108, 111, 113, 115, 119, 121, 123, 132, 135, 136]:
        if 1 <= ln <= len(door_lines):
            door_lines_sel.append((ln, door_lines[ln - 1]))
    add_code_slide(
        "Consommateur: porte",
        "DoorUnlockOnQuestsComplete.cs · event → HideDoor()",
        door_lines_sel,
        [
            "S’abonne à `QuestSystem.Instance.OnQuestsChanged`.",
            "Cache la porte uniquement si des quêtes ont été créées et qu’elles sont finies.",
            "Action: `doorSprite.enabled = false` + désactivation colliders.",
        ],
        "door_unlock",
    )

    # Slide 8: Intégration supermarché (création + progression)
    super_lines_sel = []
    for ln in [
        89, 92, 93, 94, 95, 96,
        141, 149, 151, 152, 154, 157, 168, 169,
        191, 195, 197, 200,
    ]:
        if 1 <= ln <= len(super_lines):
            super_lines_sel.append((ln, super_lines[ln - 1]))
    add_code_slide(
        "Intégration supermarché",
        "SupermarketQuestManager.cs · AddQuest + IncrementStepCount",
        super_lines_sel,
        [
            "Crée 2 quêtes: obstacle + liste de courses (avec étapes/quantités).",
            "Quand le joueur scanne un item: `IncrementStepCount()` sur l’étape correspondante.",
            "Résultat: la progression de quête remonte automatiquement vers l’UI et la porte.",
        ],
        "supermarket",
    )

    prs.save(str(OUT_PPTX))
    print(f"OK -> {OUT_PPTX}")


if __name__ == "__main__":
    main()

